from __future__ import annotations

import base64
import html as html_lib
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

ROOT = Path(__file__).resolve().parents[1]
PPTX_PATH = ROOT / "reports" / "pokemon_analysis_deck.pptx"
OUTPUT = ROOT / "reports" / "pokemon_analysis_deck.html"

EMU_PER_INCH = 914400
PX_PER_INCH = 96
PX_PER_PT = PX_PER_INCH / 72

FONT_STACK = (
    '"Apple SD Gothic Neo", "AppleGothic", "Pretendard", "Noto Sans KR", '
    '"Malgun Gothic", -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif'
)

CSS = r"""
:root {
  --yellow: #FCA311;
  --red: #E63946;
  --ink: #E9EDF5;
  --muted: #94A0B8;
  --border: rgba(255, 255, 255, .10);
  --panel: #0E1626;
}
* { box-sizing: border-box; }
html, body { height: 100%; }
body {
  margin: 0;
  background:
    radial-gradient(1200px 700px at 50% -10%, rgba(42, 157, 143, .14), transparent 60%),
    radial-gradient(900px 600px at 88% 112%, rgba(252, 163, 17, .09), transparent 55%),
    #05080F;
  color: var(--ink);
  font-family: __FONT_STACK__;
  overflow: hidden;
  -webkit-font-smoothing: antialiased;
  text-rendering: optimizeLegibility;
}
button { font-family: inherit; }
.app { height: 100vh; height: 100dvh; display: flex; flex-direction: column; }
.topbar {
  position: relative;
  z-index: 30;
  flex: 0 0 56px;
  height: 56px;
  display: flex;
  align-items: center;
  gap: 16px;
  padding: 0 20px;
  border-bottom: 1px solid rgba(255, 255, 255, .06);
  background: linear-gradient(180deg, rgba(12, 18, 32, .92), rgba(12, 18, 32, .55));
  -webkit-backdrop-filter: blur(14px);
  backdrop-filter: blur(14px);
  transition: opacity .25s ease;
}
.brand {
  font-size: 10px;
  letter-spacing: .16em;
  font-weight: 800;
  color: var(--muted);
  white-space: nowrap;
}
.brand b { color: var(--yellow); }
.doc-title {
  flex: 1;
  min-width: 0;
  text-align: center;
  font-size: 12px;
  color: var(--muted);
  white-space: nowrap;
  overflow: hidden;
  text-overflow: ellipsis;
  padding: 0 8px;
}
.top-actions { display: flex; align-items: center; gap: 6px; }
.btn {
  appearance: none;
  border: 1px solid var(--border);
  background: rgba(255, 255, 255, .04);
  color: #D7DEEB;
  font-size: 11.5px;
  font-weight: 600;
  border-radius: 8px;
  padding: 7px 10px;
  cursor: pointer;
  white-space: nowrap;
  transition: background .15s, border-color .15s, transform .15s, opacity .15s;
}
.btn:hover { background: rgba(255, 255, 255, .10); border-color: rgba(255, 255, 255, .24); }
.btn[aria-pressed="true"], .btn[aria-expanded="true"] { background: rgba(252, 163, 17, .16); border-color: rgba(252, 163, 17, .48); color: #FFF; }
.btn:active { transform: translateY(1px); }
.btn:focus-visible { outline: 2px solid var(--yellow); outline-offset: 2px; }
.btn .key {
  display: inline-block;
  margin-left: 6px;
  padding: 1px 4px;
  border: 1px solid rgba(255, 255, 255, .16);
  border-radius: 4px;
  font-size: 9px;
  color: var(--muted);
}
#stage {
  position: relative;
  flex: 1;
  min-height: 0;
  overflow: hidden;
}
#deck {
  position: absolute;
  left: 50%;
  top: 50%;
  width: 1280px;
  height: 720px;
  transform-origin: center center;
  will-change: transform;
}
.slide {
  position: absolute;
  inset: 0;
  overflow: hidden;
  background: #fff;
  border-radius: 6px;
  opacity: 0;
  pointer-events: none;
  transition: opacity .28s ease;
  box-shadow: 0 30px 80px rgba(0, 0, 0, .55), 0 0 0 1px rgba(255, 255, 255, .06);
}
.slide.active { opacity: 1; pointer-events: auto; z-index: 2; }
.text { text-rendering: geometricPrecision; }
.bottombar {
  position: relative;
  z-index: 30;
  flex: 0 0 58px;
  height: 58px;
  display: flex;
  align-items: center;
  justify-content: center;
  transition: opacity .25s ease;
}
.nav-pill {
  display: flex;
  align-items: center;
  gap: 10px;
  padding: 6px 8px;
  border: 1px solid var(--border);
  border-radius: 999px;
  background: rgba(10, 15, 26, .82);
  box-shadow: 0 12px 40px rgba(0, 0, 0, .35);
  -webkit-backdrop-filter: blur(14px);
  backdrop-filter: blur(14px);
}
.nav-btn {
  width: 30px;
  height: 30px;
  display: inline-flex;
  align-items: center;
  justify-content: center;
  border: 0;
  border-radius: 50%;
  background: transparent;
  color: #C9D3E3;
  font-size: 16px;
  cursor: pointer;
  transition: background .15s, color .15s;
}
.nav-btn:hover { background: rgba(255, 255, 255, .12); color: #fff; }
.nav-btn:focus-visible { outline: 2px solid var(--yellow); outline-offset: 1px; }
.counter {
  min-width: 58px;
  text-align: center;
  font-size: 11px;
  letter-spacing: .08em;
  color: #C1CBDB;
  font-variant-numeric: tabular-nums;
}
.progressbar {
  position: fixed;
  left: 0;
  right: 0;
  bottom: 0;
  height: 2px;
  background: rgba(255, 255, 255, .06);
  z-index: 35;
}
.progressbar span {
  display: block;
  width: 0;
  height: 100%;
  background: linear-gradient(90deg, var(--yellow), #ff8a3d);
  transition: width .25s ease;
}
body.idle .topbar, body.idle .bottombar { opacity: .22; }
body.idle .topbar:hover, body.idle .bottombar:hover { opacity: 1; }
body.notes-open #stage { opacity: .55; transition: opacity .25s ease; }
.notes-panel {
  position: fixed;
  left: 50%;
  bottom: 72px;
  z-index: 40;
  width: min(920px, calc(100vw - 40px));
  max-height: 34vh;
  display: flex;
  flex-direction: column;
  border: 1px solid rgba(255, 255, 255, .12);
  border-radius: 14px;
  background: rgba(14, 22, 38, .97);
  box-shadow: 0 24px 80px rgba(0, 0, 0, .55);
  -webkit-backdrop-filter: blur(18px);
  backdrop-filter: blur(18px);
  transform: translate(-50%, calc(100% + 90px));
  transition: transform .28s ease;
}
.notes-panel.open { transform: translate(-50%, 0); }
.notes-head {
  display: flex;
  align-items: center;
  justify-content: space-between;
  gap: 12px;
  padding: 13px 18px 10px;
  border-bottom: 1px solid rgba(255, 255, 255, .07);
  font-size: 11px;
  font-weight: 700;
  letter-spacing: .1em;
  color: var(--yellow);
}
.notes-head .notes-index { color: var(--muted); margin-left: 6px; letter-spacing: .04em; }
.notes-close {
  border: 0;
  background: transparent;
  color: var(--muted);
  font-size: 18px;
  line-height: 1;
  cursor: pointer;
  padding: 0 2px;
}
.notes-close:hover { color: #fff; }
.notes-body {
  padding: 14px 20px 20px;
  overflow: auto;
  color: #C8D2E2;
  font-size: 13.5px;
  line-height: 1.75;
  white-space: pre-wrap;
}
.overlay {
  position: fixed;
  inset: 0;
  z-index: 50;
  display: flex;
  align-items: center;
  justify-content: center;
  padding: 24px;
  background: rgba(4, 7, 12, .74);
  -webkit-backdrop-filter: blur(10px);
  backdrop-filter: blur(10px);
  opacity: 0;
  pointer-events: none;
  transition: opacity .2s ease;
}
.overlay.open { opacity: 1; pointer-events: auto; }
.help-card {
  width: min(580px, 100%);
  border: 1px solid rgba(255, 255, 255, .12);
  border-radius: 16px;
  background: #0D1524;
  box-shadow: 0 30px 90px rgba(0, 0, 0, .6);
  padding: 26px 28px 22px;
}
.help-card h2 { margin: 0 0 6px; font-size: 20px; letter-spacing: -.01em; }
.help-card p { margin: 0 0 18px; color: var(--muted); font-size: 12.5px; line-height: 1.6; }
.shortcuts { display: grid; grid-template-columns: 1fr 1fr; gap: 8px 18px; }
.shortcut {
  display: flex;
  align-items: center;
  justify-content: space-between;
  gap: 10px;
  padding: 7px 0;
  border-bottom: 1px solid rgba(255, 255, 255, .06);
  font-size: 12px;
  color: #C4CEE0;
}
kbd {
  display: inline-block;
  min-width: 22px;
  padding: 2px 6px;
  border: 1px solid rgba(255, 255, 255, .18);
  border-bottom-width: 2px;
  border-radius: 5px;
  background: rgba(255, 255, 255, .05);
  color: #E9EDF5;
  font-family: inherit;
  font-size: 10.5px;
  text-align: center;
}
.help-foot {
  display: flex;
  align-items: center;
  justify-content: space-between;
  gap: 12px;
  margin-top: 20px;
  color: var(--muted);
  font-size: 11px;
}
.help-close {
  border: 1px solid var(--border);
  background: rgba(255, 255, 255, .05);
  color: #E9EDF5;
  border-radius: 8px;
  padding: 7px 12px;
  font-size: 11.5px;
  cursor: pointer;
}
@media (max-width: 900px) {
  .topbar { padding: 0 12px; gap: 10px; }
  .brand { display: none; }
  .btn .key { display: none; }
  .doc-title { text-align: left; font-size: 11px; }
  .top-actions .btn[data-action="print"] { display: none; }
  .shortcuts { grid-template-columns: 1fr; }
}
@media (prefers-reduced-motion: reduce) {
  * { transition: none !important; animation: none !important; }
}
@media print {
  @page { size: 13.333in 7.5in; margin: 0; }
  html, body { height: auto; background: #fff; overflow: visible; }
  .app { display: block; height: auto; }
  .topbar, .bottombar, .progressbar, .notes-panel, .overlay { display: none !important; }
  #stage { display: block; position: static; height: auto; overflow: visible; opacity: 1 !important; }
  #deck { position: static; width: auto; height: auto; transform: none !important; }
  .slide {
    position: relative !important;
    inset: auto !important;
    display: block !important;
    width: 1280px;
    height: 720px;
    opacity: 1 !important;
    pointer-events: auto !important;
    border-radius: 0;
    box-shadow: none;
    page-break-after: always;
    break-after: page;
  }
  .slide:last-child { page-break-after: auto; break-after: auto; }
  .slide, .slide * {
    -webkit-print-color-adjust: exact !important;
    print-color-adjust: exact !important;
  }
}
""".replace("__FONT_STACK__", FONT_STACK)

JS = r"""
(() => {
  const NOTES = __NOTES__;
  const TITLES = __TITLES__;
  const slides = Array.from(document.querySelectorAll('.slide'));
  const deck = document.getElementById('deck');
  const stage = document.getElementById('stage');
  const counter = document.getElementById('counter');
  const progress = document.getElementById('progress');
  const titleEl = document.getElementById('slide-title');
  const notesPanel = document.getElementById('notes-panel');
  const notesText = document.getElementById('notes-text');
  const notesIndex = document.getElementById('notes-index');
  const helpOverlay = document.getElementById('help');
  const total = slides.length;

  let index = 0;
  let notesOpen = false;
  let helpOpen = false;
  let idleTimer = null;

  const initial = parseInt(location.hash.slice(1), 10);
  if (Number.isFinite(initial) && initial >= 1 && initial <= total) {
    index = initial - 1;
  }

  function render() {
    slides.forEach((slide, i) => {
      const active = i === index;
      slide.classList.toggle('active', active);
      slide.setAttribute('aria-hidden', String(!active));
      if (active) {
        slide.removeAttribute('inert');
      } else {
        slide.setAttribute('inert', '');
      }
    });
    counter.textContent = String(index + 1).padStart(2, '0') + ' / ' + String(total).padStart(2, '0');
    progress.style.width = (((index + 1) / total) * 100).toFixed(2) + '%';
    titleEl.textContent = TITLES[index] || '';
    notesText.textContent = NOTES[index] || '발표자 노트가 없습니다.';
    notesIndex.textContent = String(index + 1).padStart(2, '0');
    try {
      history.replaceState(null, '', '#' + (index + 1));
    } catch (error) {}
  }

  function go(next) {
    index = Math.max(0, Math.min(total - 1, next));
    render();
    wake();
  }

  function next() { go(index + 1); }
  function prev() { go(index - 1); }

  function fit() {
    if (!stage) return;
    const padding = window.innerWidth < 700 ? 10 : 20;
    const availableWidth = Math.max(1, stage.clientWidth - padding * 2);
    const availableHeight = Math.max(1, stage.clientHeight - padding * 2);
    const scale = Math.min(availableWidth / 1280, availableHeight / 720);
    deck.style.transform = 'translate(-50%, -50%) scale(' + scale.toFixed(4) + ')';
  }

  function setNotes(open) {
    const focusInside = document.activeElement && document.activeElement.closest('.notes-panel');
    notesOpen = open;
    notesPanel.classList.toggle('open', open);
    notesPanel.setAttribute('aria-hidden', String(!open));
    document.body.classList.toggle('notes-open', open);
    notesPanel.toggleAttribute('inert', !open);
    const notesButton = document.querySelector('[data-action="notes"]');
    if (notesButton) {
      notesButton.setAttribute('aria-pressed', String(open));
      notesButton.setAttribute('aria-expanded', String(open));
    }
    if (open) {
      notesPanel.querySelector('.notes-body').scrollTop = 0;
    } else if (focusInside && notesButton) {
      notesButton.focus({ preventScroll: true });
    }
    wake();
  }

  function setHelp(open) {
    const focusInside = document.activeElement && document.activeElement.closest('.overlay');
    helpOpen = open;
    helpOverlay.classList.toggle('open', open);
    helpOverlay.setAttribute('aria-hidden', String(!open));
    helpOverlay.toggleAttribute('inert', !open);
    const helpButton = document.querySelector('[data-action="help"]');
    if (helpButton) helpButton.setAttribute('aria-expanded', String(open));
    if (open) {
      const first = helpOverlay.querySelector('button');
      if (first) setTimeout(() => first.focus(), 30);
    } else if (focusInside && helpButton) {
      helpButton.focus({ preventScroll: true });
    }
    wake();
  }

  function toggleFullscreen() {
    if (!document.fullscreenElement) {
      document.documentElement.requestFullscreen().catch(() => {});
    } else {
      document.exitFullscreen().catch(() => {});
    }
  }

  function printDeck() {
    setNotes(false);
    setHelp(false);
    setTimeout(() => window.print(), 120);
  }

  function wake() {
    document.body.classList.remove('idle');
    clearTimeout(idleTimer);
    if (!notesOpen && !helpOpen) {
      idleTimer = setTimeout(() => document.body.classList.add('idle'), 3000);
    }
  }

  document.addEventListener('click', (event) => {
    const button = event.target.closest('[data-action]');
    if (!button) return;
    const action = button.dataset.action;
    if (action === 'next') next();
    if (action === 'prev') prev();
    if (action === 'notes') setNotes(!notesOpen);
    if (action === 'help') setHelp(!helpOpen);
    if (action === 'fullscreen') toggleFullscreen();
    if (action === 'print') printDeck();
    if (action === 'close-notes') setNotes(false);
    if (action === 'close-help') setHelp(false);
  });

  helpOverlay.addEventListener('click', (event) => {
    if (event.target === helpOverlay) setHelp(false);
  });

  document.addEventListener('keydown', (event) => {
    if (event.metaKey || event.ctrlKey || event.altKey) return;
    const key = event.key;
    if (key === 'Escape') {
      if (helpOpen) setHelp(false);
      else if (notesOpen) setNotes(false);
      return;
    }
    if (helpOpen) {
      if (key === '?' || key === '/') setHelp(false);
      return;
    }
    if (document.activeElement && document.activeElement.closest('.notes-panel')) {
      if (key === 'ArrowUp' || key === 'ArrowDown' || key === 'PageUp' || key === 'PageDown') return;
    }
    if (key === 'ArrowRight' || key === 'PageDown' || key === ' ') {
      event.preventDefault();
      next();
    } else if (key === 'ArrowLeft' || key === 'PageUp') {
      event.preventDefault();
      prev();
    } else if (key === 'Home') {
      event.preventDefault();
      go(0);
    } else if (key === 'End') {
      event.preventDefault();
      go(total - 1);
    } else if (key === 'n' || key === 'N') {
      setNotes(!notesOpen);
    } else if (key === 'f' || key === 'F') {
      toggleFullscreen();
    } else if (key === 'p' || key === 'P') {
      printDeck();
    } else if (key === '?' || key === '/') {
      setHelp(true);
    }
  });

  window.addEventListener('hashchange', () => {
    const value = parseInt(location.hash.slice(1), 10);
    if (Number.isFinite(value) && value >= 1 && value <= total && value - 1 !== index) {
      go(value - 1);
    }
  });

  let touchX = null;
  stage.addEventListener('touchstart', (event) => {
    touchX = event.touches[0].clientX;
  }, { passive: true });
  stage.addEventListener('touchend', (event) => {
    if (touchX === null) return;
    const delta = event.changedTouches[0].clientX - touchX;
    if (Math.abs(delta) > 48) {
      if (delta < 0) next();
      else prev();
    }
    touchX = null;
  }, { passive: true });

  ['mousemove', 'mousedown', 'touchstart', 'keydown', 'wheel'].forEach((name) => {
    window.addEventListener(name, wake, { passive: true });
  });

  window.addEventListener('resize', fit);
  window.addEventListener('load', fit);
  document.fonts.ready.then(fit).catch(() => {});
  if ('ResizeObserver' in window) {
    new ResizeObserver(fit).observe(stage);
  }

  render();
  fit();
  wake();
})();
"""


def px(value: int | float) -> float:
    return round(value / EMU_PER_INCH * PX_PER_INCH, 2)


def color_hex(color_format, default: str = "#000000") -> str:
    try:
        if color_format.type == 1:
            if hasattr(color_format, "rgb") and color_format.rgb is not None:
                return f"#{color_format.rgb}"
            return f"#{color_format.fore_color.rgb}"
    except Exception:
        pass
    return default


def slide_background(slide) -> str:
    try:
        fill = slide.background.fill
        if fill.type == 1:
            return color_hex(fill, "#FFFFFF")
    except Exception:
        pass
    return "#FFFFFF"


def shape_fill(shape) -> str:
    try:
        if shape.fill.type == 1:
            return color_hex(shape.fill, "transparent")
    except Exception:
        pass
    return "transparent"


def shape_border(shape) -> str:
    try:
        line_fill = shape.line.fill
        if line_fill.type == 1:
            rgb = color_hex(line_fill, "#000000")
            width = shape.line.width.pt if shape.line.width else 1
            return f"border:{width * PX_PER_PT:.2f}px solid {rgb};"
    except Exception:
        pass
    return ""


def shape_radius(shape, width: float, height: float) -> str:
    try:
        auto_type = shape.auto_shape_type
        if auto_type == MSO_SHAPE.OVAL:
            return "50%"
        if auto_type == MSO_SHAPE.ROUNDED_RECTANGLE:
            radius = shape.adjustments[0] * min(width, height)
            return f"{radius:.1f}px"
    except Exception:
        pass
    if "oval" in shape.name.lower():
        return "50%"
    return "0"


def text_html(shape) -> str:
    frame = shape.text_frame
    left = px(shape.left)
    top = px(shape.top)
    width = px(shape.width)
    height = px(shape.height)
    margin_left = px(frame.margin_left or 0)
    margin_right = px(frame.margin_right or 0)
    margin_top = px(frame.margin_top or 0)
    margin_bottom = px(frame.margin_bottom or 0)
    anchor = frame.vertical_anchor
    justify = "flex-start"
    try:
        if anchor is not None and int(anchor) == int(MSO_ANCHOR.MIDDLE):
            justify = "center"
        elif anchor is not None and int(anchor) == int(MSO_ANCHOR.BOTTOM):
            justify = "flex-end"
    except Exception:
        pass

    paragraphs = []
    for paragraph in frame.paragraphs:
        if not paragraph.runs:
            continue
        align = "left"
        if paragraph.alignment == PP_ALIGN.CENTER:
            align = "center"
        elif paragraph.alignment == PP_ALIGN.RIGHT:
            align = "right"

        spacing = paragraph.line_spacing
        if isinstance(spacing, float):
            line_height = f"{spacing * 1.2:.3f}"
        elif spacing is not None:
            line_height = f"{spacing.pt * PX_PER_PT:.2f}px"
        else:
            line_height = "1.2"
        after = paragraph.space_after.pt * PX_PER_PT if paragraph.space_after is not None else 0

        spans = []
        for run in paragraph.runs:
            size = run.font.size.pt if run.font.size else 14
            weight = "700" if run.font.bold else "400"
            color = color_hex(run.font.color, "#1B1F2A")
            spans.append(
                f'<span style="font-size:{size * PX_PER_PT:.2f}px;font-weight:{weight};color:{color};">'
                f"{html_lib.escape(run.text)}</span>"
            )
        paragraphs.append(
            f'<div style="margin:0 0 {after:.2f}px 0;text-align:{align};line-height:{line_height};">'
            f'{"".join(spans)}</div>'
        )

    return (
        f'<div class="text" style="position:absolute;left:{left}px;top:{top}px;width:{width}px;'
        f"height:{height}px;padding:{margin_top}px {margin_right}px {margin_bottom}px {margin_left}px;"
        f'display:flex;flex-direction:column;justify-content:{justify};box-sizing:border-box;'
        f'white-space:pre-wrap;overflow:visible;">{"".join(paragraphs)}</div>'
    )


def render_shape(shape) -> str:
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        left, top = px(shape.left), px(shape.top)
        width, height = px(shape.width), px(shape.height)
        blob = base64.b64encode(shape.image.blob).decode("ascii")
        mime = shape.image.content_type or "image/png"
        return (
            f'<img src="data:{mime};base64,{blob}" alt="" draggable="false" '
            f'style="position:absolute;left:{left}px;top:{top}px;width:{width}px;height:{height}px;">'
        )

    parts = []
    if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        left, top = px(shape.left), px(shape.top)
        width, height = px(shape.width), px(shape.height)
        fill = shape_fill(shape)
        border = shape_border(shape)
        radius = shape_radius(shape, width, height)
        rotation = f"transform:rotate({shape.rotation}deg);" if getattr(shape, "rotation", 0) else ""
        parts.append(
            f'<div style="position:absolute;left:{left}px;top:{top}px;width:{width}px;height:{height}px;'
            f'background:{fill};{border}border-radius:{radius};{rotation}"></div>'
        )

    if shape.has_text_frame and shape.text_frame.text.strip():
        parts.append(text_html(shape))
    return "".join(parts)


def extract_title(slide) -> str:
    candidates = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = " ".join(p.text for p in shape.text_frame.paragraphs).strip()
        if not text:
            continue
        sizes = [
            run.font.size.pt
            for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs
            if run.font.size
        ]
        if not sizes:
            continue
        candidates.append((max(sizes), shape.top / EMU_PER_INCH, text))
    if not candidates:
        return "포켓몬 데이터 분석"
    preferred = [item for item in candidates if item[0] >= 20 and item[1] < 1.5]
    pool = preferred or candidates
    return max(pool, key=lambda item: (item[0], -item[1]))[2]


def render_slide(slide, index: int) -> tuple[str, str, str]:
    title = extract_title(slide)
    background = slide_background(slide)
    notes = ""
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
    parts = [
        f'<section class="slide" role="group" aria-label="{html_lib.escape(title)}" '
        f'style="background:{background};">'
    ]
    for shape in slide.shapes:
        parts.append(render_shape(shape))
    parts.append("</section>")
    return "".join(parts), title, notes


def main() -> None:
    prs = Presentation(PPTX_PATH)
    rendered = [render_slide(slide, i) for i, slide in enumerate(prs.slides, start=1)]
    slides_html = "\n".join(item[0] for item in rendered)
    titles = [item[1] for item in rendered]
    notes = [item[2] for item in rendered]

    notes_json = json.dumps(notes, ensure_ascii=False).replace("</", "<\\/")
    titles_json = json.dumps(titles, ensure_ascii=False).replace("</", "<\\/")
    script = JS.replace("__NOTES__", notes_json).replace("__TITLES__", titles_json)

    document = f"""<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="color-scheme" content="dark">
<title>포켓몬 데이터 분석 · 2026</title>
<style>
{CSS}
</style>
</head>
<body>
<div class="app">
  <header class="topbar">
    <div class="brand">POKÉMON <b>·</b> DATA ANALYSIS</div>
    <div class="doc-title" id="slide-title"></div>
    <div class="top-actions">
      <button class="btn" type="button" data-action="notes" aria-pressed="false" aria-expanded="false">노트<span class="key">N</span></button>
      <button class="btn" type="button" data-action="fullscreen">전체화면<span class="key">F</span></button>
      <button class="btn" type="button" data-action="print">PDF<span class="key">P</span></button>
      <button class="btn" type="button" data-action="help" aria-label="도움말" aria-expanded="false">도움말</button>
    </div>
  </header>
  <main id="stage">
    <div id="deck">
{slides_html}
    </div>
  </main>
  <footer class="bottombar">
    <div class="nav-pill">
      <button class="nav-btn" type="button" data-action="prev" aria-label="이전 슬라이드">←</button>
      <div class="counter" id="counter">01 / {len(rendered)}</div>
      <button class="nav-btn" type="button" data-action="next" aria-label="다음 슬라이드">→</button>
    </div>
  </footer>
</div>
<div class="progressbar"><span id="progress"></span></div>
<aside class="notes-panel" id="notes-panel" aria-hidden="true" inert>
  <div class="notes-head">
    <div>발표자 노트 <span class="notes-index" id="notes-index">01</span></div>
    <button class="notes-close" type="button" data-action="close-notes" aria-label="노트 닫기">×</button>
  </div>
  <div class="notes-body" id="notes-text"></div>
</aside>
<div class="overlay" id="help" aria-hidden="true" inert>
  <div class="help-card" role="dialog" aria-modal="true" aria-labelledby="help-title">
    <h2 id="help-title">키보드 단축키</h2>
    <p>화면을 클릭한 뒤 키보드로 이동하세요. 인쇄(P)에서 PDF로 저장하면 슬라이드 원본 비율 그대로 출력됩니다.</p>
    <div class="shortcuts">
      <div class="shortcut"><span>다음 슬라이드</span><span><kbd>→</kbd> <kbd>Space</kbd></span></div>
      <div class="shortcut"><span>이전 슬라이드</span><span><kbd>←</kbd></span></div>
      <div class="shortcut"><span>처음 / 끝</span><span><kbd>Home</kbd> <kbd>End</kbd></span></div>
      <div class="shortcut"><span>발표자 노트</span><span><kbd>N</kbd></span></div>
      <div class="shortcut"><span>전체화면</span><span><kbd>F</kbd></span></div>
      <div class="shortcut"><span>인쇄 / PDF</span><span><kbd>P</kbd></span></div>
      <div class="shortcut"><span>도움말 닫기</span><span><kbd>Esc</kbd></span></div>
    </div>
    <div class="help-foot">
      <span>발표자 노트는 슬라이드마다 포함되어 있습니다.</span>
      <button class="help-close" type="button" data-action="close-help">닫기</button>
    </div>
  </div>
</div>
<script>
{script}
</script>
</body>
</html>
"""
    OUTPUT.write_text(document, encoding="utf-8")
    print(f"Saved {OUTPUT} ({len(rendered)} slides, {OUTPUT.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
