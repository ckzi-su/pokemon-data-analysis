from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from matplotlib.patches import Patch
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data" / "raw" / "Pokemon.csv"
ASSET_DIR = ROOT / "reports" / "deck_assets"
OUTPUT = ROOT / "reports" / "pokemon_analysis_deck.pptx"

FONT = "Apple SD Gothic Neo"
NAVY_HEX = "#14213D"
YELLOW_HEX = "#FCA311"
RED_HEX = "#E63946"
TEAL_HEX = "#2A9D8F"
GRAY_HEX = "#5B6472"
INK_HEX = "#1B1F2A"
LIGHT_HEX = "#F1F3F7"
BORDER_HEX = "#DDE3EC"

NAVY = RGBColor(0x14, 0x21, 0x3D)
YELLOW = RGBColor(0xFC, 0xA3, 0x11)
RED = RGBColor(0xE6, 0x39, 0x46)
INK = RGBColor(0x1B, 0x1F, 0x2A)
GRAY = RGBColor(0x5B, 0x64, 0x72)
MUTED = RGBColor(0x8A, 0x93, 0xA3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF1, 0xF3, 0xF7)
BORDER = RGBColor(0xDD, 0xE3, 0xEC)
PALE_YELLOW = RGBColor(0xFF, 0xF4, 0xD6)
PALE_RED = RGBColor(0xFD, 0xEC, 0xED)

STATS = ["hp", "attack", "defense", "sp_atk", "sp_def", "speed"]
STAT_LABELS = ["HP", "공격", "방어", "특공", "특방", "스피드"]
DISPLAY_NAMES = {
    "MewtwoMega Mewtwo X": "메가뮤츠 X",
    "MewtwoMega Mewtwo Y": "메가뮤츠 Y",
    "RayquazaMega Rayquaza": "메가레쿠쟈",
    "KyogrePrimal Kyogre": "원시가이오가",
    "GroudonPrimal Groudon": "원시그란돈",
    "Arceus": "아르세우스",
    "TyranitarMega Tyranitar": "메가마기라스",
    "SalamenceMega Salamence": "메가보만다",
    "MetagrossMega Metagross": "메가메타그로스",
    "LatiasMega Latias": "메가라티아스",
}


def load_data() -> pd.DataFrame:
    df = pd.read_csv(DATA_PATH)
    df = df.rename(
        columns={
            "#": "id",
            "Name": "name",
            "Type 1": "type_1",
            "Type 2": "type_2",
            "Total": "total",
            "HP": "hp",
            "Attack": "attack",
            "Defense": "defense",
            "Sp. Atk": "sp_atk",
            "Sp. Def": "sp_def",
            "Speed": "speed",
            "Generation": "generation",
            "Legendary": "legendary",
        }
    )
    df["type_2"] = df["type_2"].fillna("None")
    df["legendary"] = df["legendary"].astype(bool)
    return df


def setup_matplotlib() -> None:
    plt.rcParams.update(
        {
            "font.family": FONT,
            "axes.unicode_minus": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "axes.edgecolor": BORDER_HEX,
            "axes.labelcolor": INK_HEX,
            "xtick.color": GRAY_HEX,
            "ytick.color": GRAY_HEX,
            "text.color": INK_HEX,
            "axes.grid": True,
            "grid.color": "#ECEFF4",
            "grid.linewidth": 0.8,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "font.size": 11,
            "legend.frameon": False,
        }
    )


def style_axes(ax: plt.Axes) -> None:
    ax.tick_params(length=0)
    ax.grid(axis="x", color="#ECEFF4")
    ax.grid(axis="y", visible=False)
    ax.set_axisbelow(True)


def save_figure(fig: plt.Figure, name: str) -> Path:
    path = ASSET_DIR / name
    fig.savefig(path, dpi=200, bbox_inches="tight", pad_inches=0.12, facecolor="white")
    plt.close(fig)
    return path


def chart_distribution(df: pd.DataFrame) -> Path:
    fig, ax = plt.subplots(figsize=(7.6, 4.8))
    sns.histplot(df["total"], bins=30, color=NAVY_HEX, edgecolor="white", ax=ax)
    mean_total = df["total"].mean()
    median_total = df["total"].median()
    ax.axvline(mean_total, color=TEAL_HEX, linewidth=2, label=f"평균 {mean_total:.1f}")
    ax.axvline(median_total, color=RED_HEX, linewidth=2, linestyle="--", label=f"중앙값 {median_total:.0f}")
    ax.set_xlabel("종족값 총합", fontsize=12)
    ax.set_ylabel("포켓몬 수", fontsize=12)
    ax.set_title("종족값 총합 분포", loc="left", fontsize=15, fontweight="bold", pad=12)
    ax.legend(loc="upper right", fontsize=10.5, handlelength=1.6)
    style_axes(ax)
    return save_figure(fig, "dist_total.png")


def chart_top10(df: pd.DataFrame) -> Path:
    top = df.nlargest(10, "total").iloc[::-1]
    labels = [DISPLAY_NAMES.get(name, name) for name in top["name"]]
    colors = [YELLOW_HEX if flag else NAVY_HEX for flag in top["legendary"]]
    fig, ax = plt.subplots(figsize=(7.6, 5.2))
    bars = ax.barh(labels, top["total"], color=colors, height=0.68)
    for bar, value in zip(bars, top["total"]):
        ax.text(value + 6, bar.get_y() + bar.get_height() / 2, f"{value}", va="center", fontsize=10.5, color=GRAY_HEX)
    ax.set_xlim(0, 850)
    ax.set_xlabel("종족값 총합", fontsize=12)
    ax.set_title("종족값 상위 10", loc="left", fontsize=15, fontweight="bold", pad=12)
    handles = [Patch(color=YELLOW_HEX, label="전설"), Patch(color=NAVY_HEX, label="일반")]
    ax.legend(handles=handles, loc="lower right", fontsize=10.5)
    style_axes(ax)
    ax.grid(axis="y", visible=False)
    return save_figure(fig, "top10_total.png")


def chart_type_total(df: pd.DataFrame) -> Path:
    long = df.melt(id_vars=["total"], value_vars=["type_1", "type_2"], value_name="type")
    long = long[long["type"] != "None"]
    means = long.groupby("type")["total"].mean().sort_values()
    colors = [YELLOW_HEX if name == means.index[-1] else RED_HEX if name == means.index[0] else NAVY_HEX for name in means.index]
    fig, ax = plt.subplots(figsize=(7.8, 5.4))
    bars = ax.barh(means.index, means.values, color=colors, height=0.7)
    for bar, value in zip(bars, means.values):
        ax.text(value + 4, bar.get_y() + bar.get_height() / 2, f"{value:.0f}", va="center", fontsize=9.5, color=GRAY_HEX)
    ax.set_xlim(0, 610)
    ax.set_xlabel("평균 종족값", fontsize=12)
    ax.set_title("타입별 평균 종족값", loc="left", fontsize=15, fontweight="bold", pad=12)
    style_axes(ax)
    ax.grid(axis="y", visible=False)
    return save_figure(fig, "type_total.png")


def chart_legendary(df: pd.DataFrame) -> Path:
    regular = df[~df["legendary"]][STATS].mean()
    legendary = df[df["legendary"]][STATS].mean()
    x = np.arange(len(STATS))
    width = 0.36
    fig, ax = plt.subplots(figsize=(7.4, 4.9))
    bars_regular = ax.bar(x - width / 2, regular.values, width, color="#9AA7BC", label="일반")
    bars_legendary = ax.bar(x + width / 2, legendary.values, width, color=YELLOW_HEX, label="전설")
    for bars in (bars_regular, bars_legendary):
        for bar in bars:
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + 3,
                f"{bar.get_height():.0f}",
                ha="center",
                fontsize=9,
                color=GRAY_HEX,
            )
    ax.set_xticks(x, STAT_LABELS, fontsize=11)
    ax.set_ylim(0, 140)
    ax.set_ylabel("평균 능력치", fontsize=12)
    ax.set_title("전설 vs 일반: 평균 능력치", loc="left", fontsize=15, fontweight="bold", pad=12)
    ax.legend(loc="upper left", fontsize=10.5, ncol=2)
    style_axes(ax)
    return save_figure(fig, "legendary_comparison.png")


def chart_correlation(df: pd.DataFrame) -> Path:
    columns = STATS + ["total"]
    labels = STAT_LABELS + ["총합"]
    corr = df[columns].corr()
    fig, ax = plt.subplots(figsize=(5.6, 5.1))
    sns.heatmap(
        corr,
        annot=False,
        cmap="Blues",
        vmin=0,
        vmax=1.0,
        square=True,
        linewidths=2,
        linecolor="white",
        cbar_kws={"shrink": 0.72},
        xticklabels=labels,
        yticklabels=labels,
        ax=ax,
    )
    for row in range(len(corr)):
        for col in range(len(corr)):
            value = corr.iat[row, col]
            ax.text(
                col + 0.5,
                row + 0.5,
                f"{value:.2f}",
                ha="center",
                va="center",
                fontsize=9.5,
                color="white" if value >= 0.55 else INK_HEX,
            )
    ax.grid(False)
    ax.set_title("능력치 간 상관계수", loc="left", fontsize=15, fontweight="bold", pad=12)
    ax.tick_params(axis="x", rotation=35, labelsize=10.5)
    ax.tick_params(axis="y", rotation=0, labelsize=10.5)
    return save_figure(fig, "correlation.png")


def chart_generation(df: pd.DataFrame) -> Path:
    grouped = df.groupby("generation").agg(mean_total=("total", "mean"), legends=("legendary", "sum"))
    fig, axes = plt.subplots(1, 2, figsize=(8.4, 4.7), gridspec_kw={"width_ratios": [1.25, 1]})
    bars = axes[0].bar(grouped.index, grouped["mean_total"], color=NAVY_HEX, width=0.62)
    for bar, value in zip(bars, grouped["mean_total"]):
        axes[0].text(bar.get_x() + bar.get_width() / 2, value + 2, f"{value:.0f}", ha="center", fontsize=10, color=GRAY_HEX)
    axes[0].set_ylim(380, 480)
    axes[0].set_xticks(grouped.index, [f"{g}세대" for g in grouped.index])
    axes[0].set_title("세대별 평균 종족값", loc="left", fontsize=14, fontweight="bold", pad=10)
    axes[0].grid(axis="y", color="#ECEFF4")
    axes[0].grid(axis="x", visible=False)
    axes[0].set_axisbelow(True)
    axes[0].tick_params(length=0)
    axes[0].spines["left"].set_visible(False)
    axes[0].spines["bottom"].set_visible(False)

    bars = axes[1].bar(grouped.index, grouped["legends"], color=YELLOW_HEX, width=0.62)
    for bar, value in zip(bars, grouped["legends"]):
        axes[1].text(bar.get_x() + bar.get_width() / 2, value + 0.5, f"{int(value)}", ha="center", fontsize=10, color=GRAY_HEX)
    axes[1].set_ylim(0, 22)
    axes[1].set_yticks([0, 5, 10, 15, 20])
    axes[1].set_xticks(grouped.index, [f"{g}" for g in grouped.index])
    axes[1].set_xlabel("세대", fontsize=11)
    axes[1].set_title("세대별 전설 수", loc="left", fontsize=14, fontweight="bold", pad=10)
    axes[1].grid(axis="y", color="#ECEFF4")
    axes[1].grid(axis="x", visible=False)
    axes[1].set_axisbelow(True)
    axes[1].tick_params(length=0)
    axes[1].spines["left"].set_visible(False)
    axes[1].spines["bottom"].set_visible(False)
    fig.tight_layout(w_pad=2.2)
    return save_figure(fig, "generation.png")


def style_run(run, size: float, bold: bool = False, color: RGBColor = INK) -> None:
    font = run.font
    font.name = FONT
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    latin = rpr.find(qn("a:latin"))
    existing = rpr.find(qn("a:ea"))
    if existing is None:
        ea = rpr.makeelement(qn("a:ea"), {"typeface": FONT})
        if latin is not None:
            latin.addnext(ea)
        else:
            rpr.append(ea)
    else:
        existing.set("typeface", FONT)


def add_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text,
    size: float = 16,
    bold: bool = False,
    color: RGBColor = INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing: float = 1.12,
    space_after: float = 0,
) -> object:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    lines = text if isinstance(text, (list, tuple)) else [text]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = line
        style_run(run, size, bold, color)
    return box


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor,
    line: RGBColor | None = None,
    line_width: float = 1.0,
    shape=MSO_SHAPE.RECTANGLE,
    radius: float | None = None,
) -> object:
    rect = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    rect.fill.solid()
    rect.fill.fore_color.rgb = fill
    if line is None:
        rect.line.fill.background()
    else:
        rect.line.color.rgb = line
        rect.line.width = Pt(line_width)
    rect.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        rect.adjustments[0] = radius
    return rect


def add_card(slide, x: float, y: float, w: float, h: float, fill: RGBColor = LIGHT, accent: RGBColor | None = None):
    add_rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.055)
    if accent is not None:
        add_rect(slide, x, y, 0.07, h, accent)


def add_bullets(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    items,
    size: float = 13.5,
    color: RGBColor = GRAY,
    lead_color: RGBColor = NAVY,
    space_after: float = 8,
    bullet: str = "•",
    line_spacing: float = 1.18,
) -> object:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.line_spacing = line_spacing
        paragraph.space_after = Pt(space_after)
        marker = paragraph.add_run()
        marker.text = f"{bullet}  "
        style_run(marker, size, True, YELLOW)
        if isinstance(item, (tuple, list)):
            lead, rest = item
            run = paragraph.add_run()
            run.text = lead
            style_run(run, size, True, lead_color)
            run = paragraph.add_run()
            run.text = f"  {rest}"
            style_run(run, size, False, color)
        else:
            run = paragraph.add_run()
            run.text = item
            style_run(run, size, False, color)
    return box


def add_stat_strip(slide, x: float, y: float, w: float, stats, value_size: float = 20, label_size: float = 9.5) -> None:
    add_rect(slide, x, y, w, 0.012, BORDER)
    column_w = w / len(stats)
    for index, (value, label) in enumerate(stats):
        left = x + index * column_w
        add_text(slide, left, y + 0.17, column_w - 0.08, 0.34, value, size=value_size, bold=True, color=NAVY)
        add_text(slide, left, y + 0.52, column_w - 0.08, 0.22, label, size=label_size, color=MUTED)


def add_image_fit(slide, path: Path, x: float, y: float, w: float, h: float, align: str = "center") -> None:
    image_w, image_h = Image.open(path).size
    ratio = min(w / image_w, h / image_h)
    draw_w = image_w * ratio
    draw_h = image_h * ratio
    if align == "center":
        left = x + (w - draw_w) / 2
    else:
        left = x
    top = y + (h - draw_h) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(draw_w), Inches(draw_h))


def add_footer(slide, page: int, total: int = 13) -> None:
    add_rect(slide, 0.8, 6.98, 11.73, 0.012, BORDER)
    add_text(
        slide,
        0.8,
        7.06,
        8.5,
        0.25,
        "출처: Kaggle — abcsds/pokemon  ·  분석 코드: src/analysis.py",
        size=9,
        color=MUTED,
    )
    add_text(
        slide,
        10.9,
        7.05,
        1.63,
        0.25,
        f"{page:02d} / {total}",
        size=9.5,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def new_content_slide(prs: Presentation, kicker: str, title: str, page: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = WHITE
    add_text(slide, 0.8, 0.44, 8.0, 0.28, kicker, size=11, bold=True, color=RED)
    add_text(slide, 0.8, 0.76, 11.7, 0.62, title, size=27, bold=True, color=NAVY)
    add_rect(slide, 0.8, 1.52, 0.82, 0.055, YELLOW)
    add_footer(slide, page)
    return slide


def build_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY

    add_rect(slide, 9.7, 1.2, 5.2, 5.2, NAVY, line=YELLOW, line_width=2.5, shape=MSO_SHAPE.OVAL)
    add_rect(slide, 9.7, 3.77, 5.2, 0.07, YELLOW)
    add_rect(slide, 11.8, 3.3, 1.0, 1.0, NAVY, line=YELLOW, line_width=2.5, shape=MSO_SHAPE.OVAL)

    add_text(slide, 0.9, 0.9, 8.4, 0.35, "DATA ANALYSIS  ·  2026", size=13, bold=True, color=YELLOW)
    add_text(slide, 0.9, 2.0, 8.3, 2.3, ["포켓몬 800마리,", "능력치로 읽다"], size=52, bold=True, color=WHITE, line_spacing=1.02)
    add_rect(slide, 0.9, 4.48, 1.15, 0.06, YELLOW)
    add_text(
        slide,
        0.9,
        4.8,
        7.6,
        0.9,
        ["세대 · 타입 · 전설 여부로 살펴본 종족값 구조", "Kaggle Pokemon Dataset (800 Pokémon / 13 features)"],
        size=17,
        color=RGBColor(0xC9, 0xD1, 0xDF),
        line_spacing=1.35,
    )
    add_text(slide, 0.9, 6.55, 7.5, 0.35, "2026. 10", size=12, color=RGBColor(0x9A, 0xA5, 0xB8))
    add_notes(
        slide,
        "안녕하세요. 오늘은 Kaggle의 포켓몬 데이터셋 800마리를 활용해 종족값 구조를 분석한 결과를 공유하겠습니다. "
        "타입, 세대, 전설 여부라는 세 가지 축으로 '강함'이 어떻게 구성되는지 살펴봤습니다.",
    )


def build_agenda_slide(prs: Presentation) -> None:
    slide = new_content_slide(prs, "CONTENTS", "오늘 다룰 내용", 2)
    items = [
        ("01", "데이터 개요", "800마리 · 13개 변수 · 결측치 0"),
        ("02", "핵심 발견", "숫자 4개로 보는 전체 요약"),
        ("03", "분포와 상위권", "종족값은 어떻게 퍼져 있는가"),
        ("04", "타입과 전설", "가장 강한 타입, 전설의 위상"),
        ("05", "상관과 세대", "능력치 조합과 세대별 추이"),
        ("06", "결론과 한계", "인사이트 정리 및 다음 단계"),
    ]
    for index, (number, title, description) in enumerate(items):
        column = index % 2
        row = index // 2
        x = 0.85 + column * 6.05
        y = 1.95 + row * 1.55
        add_card(slide, x, y, 5.8, 1.3)
        add_text(slide, x + 0.35, y + 0.22, 0.75, 0.6, number, size=24, bold=True, color=YELLOW)
        add_text(slide, x + 1.2, y + 0.2, 4.2, 0.4, title, size=18, bold=True, color=NAVY)
        add_text(slide, x + 1.2, y + 0.68, 4.3, 0.35, description, size=12.5, color=GRAY)
    add_notes(slide, "크게 데이터 소개, 핵심 발견, 분포와 상위권, 타입과 전설, 상관과 세대, 그리고 결론 순으로 구성했습니다.")


def build_data_slide(prs: Presentation, df: pd.DataFrame) -> None:
    slide = new_content_slide(prs, "01 · DATA", "데이터 개요와 분석 방법", 3)
    cards = [
        (f"{len(df):,}", "포켓몬 레코드", "1–6세대 전체"),
        (f"{df.shape[1]}", "변수", "타입 · 능력치 · 메타"),
        ("0", "결측치", "분석 전 정제 완료"),
    ]
    for index, (number, label, sub) in enumerate(cards):
        x = 0.8 + index * 4.0
        add_card(slide, x, 1.82, 3.73, 1.32)
        add_text(slide, x + 0.3, 1.98, 1.6, 0.55, number, size=28, bold=True, color=NAVY)
        add_text(slide, x + 1.65, 2.03, 1.9, 0.3, label, size=12, bold=True, color=GRAY)
        add_text(slide, x + 1.65, 2.38, 1.95, 0.3, sub, size=10.5, color=MUTED)

    add_card(slide, 0.8, 3.42, 5.72, 2.88)
    add_text(slide, 1.1, 3.65, 5.0, 0.35, "변수 구성", size=17, bold=True, color=NAVY)
    add_bullets(
        slide,
        1.1,
        4.12,
        5.15,
        2.0,
        [
            ("식별", "번호 · 이름"),
            ("타입", "Type 1 · Type 2 (복합 타입 51.8%)"),
            ("능력치", "HP · 공격 · 방어 · 특공 · 특방 · 스피드"),
            ("메타", "총합 · 세대 · 전설 여부"),
        ],
        size=12.5,
        space_after=7,
    )

    add_card(slide, 6.8, 3.42, 5.73, 2.88)
    add_text(slide, 7.1, 3.65, 5.0, 0.35, "분석 절차", size=17, bold=True, color=NAVY)
    add_bullets(
        slide,
        7.1,
        4.12,
        5.1,
        2.0,
        [
            ("01", "전처리 — 타입 결측 정리, Boolean 변환"),
            ("02", "기술통계 — 분포·타입·세대별 요약"),
            ("03", "그룹 비교 — 전설 vs 일반"),
            ("04", "상관 분석 — 능력치와 총합의 관계"),
        ],
        size=12.5,
        space_after=7,
    )

    add_rect(slide, 0.8, 6.48, 11.73, 0.36, PALE_YELLOW, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
    add_text(
        slide,
        1.05,
        6.54,
        11.3,
        0.26,
        "총합(Total)은 6개 능력치의 단순 합 — 원본 데이터에서 100% 일치 확인",
        size=11,
        bold=True,
        color=NAVY,
    )
    add_notes(
        slide,
        "데이터는 Kaggle의 abcsds/pokemon 데이터셋으로, 1세대부터 6세대까지 800마리, 13개 변수로 구성됩니다. "
        "결측치는 없고, 복합 타입 비율은 51.8%입니다. 총합은 6개 능력치를 그대로 더한 값임을 먼저 확인했습니다.",
    )


def build_kpi_slide(prs: Presentation, df: pd.DataFrame, dragon_mean: float, legendary_total: float, regular_total: float) -> None:
    slide = new_content_slide(prs, "02 · KEY FINDINGS", "한눈에 보는 핵심 발견", 4)
    cards = [
        (f"{len(df):,}", "포켓몬", "1–6세대 전체"),
        (f"{int(df['legendary'].sum())}", "전설 포켓몬", f"전체의 {df['legendary'].mean() * 100:.1f}%"),
        (f"{dragon_mean:.1f}", "Dragon 평균 총합", "18개 타입 중 1위"),
        (f"{legendary_total:.1f}", "전설 평균 총합", f"일반 {regular_total:.1f} 대비 +52.8%"),
    ]
    for index, (number, label, sub) in enumerate(cards):
        x = 0.8 + index * 3.0
        add_card(slide, x, 1.82, 2.72, 2.02)
        add_rect(slide, x, 1.82, 2.72, 0.07, YELLOW)
        add_text(slide, x + 0.25, 2.05, 2.25, 0.62, number, size=30, bold=True, color=NAVY)
        add_text(slide, x + 0.25, 2.75, 2.25, 0.3, label, size=12.5, bold=True, color=GRAY)
        add_text(slide, x + 0.25, 3.15, 2.3, 0.5, sub, size=10.5, color=MUTED, line_spacing=1.2)

    findings = [
        ("전설은 다른 리그", "평균 총합 637.4로 일반보다 52.8% 높고, 6개 능력치 모두 우위"),
        ("Dragon이 최강 타입", "평균 541.8로 1위, 최하위 Bug와는 162.3 차이"),
        ("상위권은 전설·메가", "상위 50위 중 56%가 전설, Top10 중 7마리"),
    ]
    for index, (title, body) in enumerate(findings):
        x = 0.8 + index * 4.0
        add_card(slide, x, 4.28, 3.73, 1.85)
        add_text(slide, x + 0.28, 4.5, 3.2, 0.35, title, size=15, bold=True, color=NAVY)
        add_text(slide, x + 0.28, 4.95, 3.18, 1.0, body, size=12, color=GRAY, line_spacing=1.25)
    add_text(
        slide,
        0.8,
        6.42,
        11.7,
        0.35,
        "→ 강함의 구조는 '전설성'과 '타입'이 크게 좌우하며, 상위권은 특수한 형태가 독식",
        size=12,
        bold=True,
        color=RED,
    )
    add_notes(
        slide,
        "핵심 숫자 네 가지입니다. 800마리 중 65마리가 전설이고, 전설의 평균 총합은 637.4로 일반보다 52.8% 높습니다. "
        "타입 중에서는 Dragon이 541.8로 가장 강했습니다.",
    )


def build_distribution_slide(prs: Presentation, chart: Path) -> None:
    slide = new_content_slide(prs, "03 · DISTRIBUTION", "종족값은 어디에 몰려 있는가", 5)
    add_image_fit(slide, chart, 0.72, 1.72, 7.7, 4.95)
    add_card(slide, 8.7, 1.85, 3.83, 4.65)
    add_text(slide, 9.0, 2.08, 3.3, 0.35, "분포 읽기", size=17, bold=True, color=NAVY)
    add_bullets(
        slide,
        9.0,
        2.58,
        3.25,
        3.7,
        [
            ("450", "중앙값 · 평균 435.1보다 높음"),
            ("54.8%", "350–550 구간에 집중"),
            ("48마리", "600 초과 · 상위 6.0%"),
            ("왼쪽 꼬리", "저종족값 포켓몬이 평균을 끌어내림"),
        ],
        size=12.5,
        space_after=13,
    )
    add_stat_strip(slide, 9.0, 5.35, 3.23, [("48", "600 초과"), ("6.0%", "상위 비율"), ("780", "최고 종족값")])
    add_notes(
        slide,
        "종족값 분포를 보면 350에서 550 사이에 절반 이상이 몰려 있습니다. 중앙값 450이 평균 435.1보다 높아, "
        "낮은 종족값 쪽 꼬리가 더 길다는 점이 특징입니다. 600을 넘는 48마리는 대부분 전설이나 메가진화입니다.",
    )


def build_top10_slide(prs: Presentation, chart: Path) -> None:
    slide = new_content_slide(prs, "04 · TOP POKÉMON", "상위권은 누가 차지했나", 6)
    add_image_fit(slide, chart, 0.72, 1.7, 7.65, 5.05)
    add_card(slide, 8.65, 1.85, 3.88, 4.65)
    add_text(slide, 8.95, 2.08, 3.3, 0.35, "Top 10 분석", size=17, bold=True, color=NAVY)
    add_bullets(
        slide,
        8.95,
        2.58,
        3.3,
        3.75,
        [
            ("780", "메가뮤츠 X·Y, 메가레쿠쟈 공동 1위"),
            ("7 / 10", "상위 10마리 중 전설"),
            ("56%", "상위 50위 중 전설 비율"),
            ("메가진화", "일반 3마리는 모두 메가진화 형태"),
        ],
        size=12.5,
        space_after=13,
    )
    add_stat_strip(slide, 8.95, 5.35, 3.28, [("7", "Top10 전설"), ("3", "Top10 일반"), ("56%", "Top50 전설")])
    add_notes(
        slide,
        "상위 10위 안에는 전설이 7마리, 일반이 3마리입니다. 일반 3마리도 모두 메가마기라스, 메가보만다, 메가메타그로스처럼 "
        "메가진화 형태입니다. 데이터에 메가진화와 원시회귀가 포함되어 있다는 점은 해석 시 주의해야 합니다.",
    )


def build_type_slide(prs: Presentation, chart: Path) -> None:
    slide = new_content_slide(prs, "05 · TYPE", "어떤 타입이 가장 강한가", 7)
    add_image_fit(slide, chart, 0.7, 1.72, 7.65, 5.0)
    add_card(slide, 8.6, 1.85, 3.93, 4.65)
    add_text(slide, 8.9, 2.08, 3.35, 0.35, "타입별 비교", size=17, bold=True, color=NAVY)
    add_bullets(
        slide,
        8.9,
        2.58,
        3.35,
        3.8,
        [
            ("541.8", "Dragon — 18개 타입 중 1위"),
            ("486.6", "Steel 2위 · Psychic 477.1 3위"),
            ("379.5", "Bug 최하위 · Poison 397.7"),
            ("162.3", "1위와 최하위의 평균 격차"),
            ("51.8%", "복합 타입은 양쪽에 중복 집계"),
        ],
        size=12.5,
        space_after=10,
    )
    add_stat_strip(slide, 8.9, 5.35, 3.33, [("18", "타입 수"), ("162.3", "1–18위 격차"), ("126", "Water 개체")])
    add_notes(
        slide,
        "타입별 평균 종족값에서는 Dragon이 541.8로 가장 높았고, Steel과 Psychic이 뒤를 이었습니다. "
        "Bug는 379.5로 가장 낮아 1위와 162.3 차이가 났습니다. 복합 타입은 두 타입 모두에 집계됐고, "
        "타입별 표본 수가 38에서 126까지 차이 나므로 평균 해석에 주의해야 합니다.",
    )


def build_legendary_slide(prs: Presentation, chart: Path, gaps: pd.Series) -> None:
    slide = new_content_slide(prs, "06 · LEGENDARY", "전설은 정말 다른 리그인가", 8)
    add_image_fit(slide, chart, 0.72, 1.78, 7.4, 4.85)
    add_card(slide, 8.55, 1.85, 3.98, 4.65)
    add_text(slide, 8.85, 2.08, 3.4, 0.35, "일반 vs 전설", size=17, bold=True, color=NAVY)
    add_bullets(
        slide,
        8.85,
        2.58,
        3.4,
        3.8,
        [
            ("+52.8%", "전설 평균 총합 637.4 vs 일반 417.2"),
            ("6 / 6", "모든 능력치에서 전설 우위"),
            (f"+{gaps['sp_atk']:.1f}", "최대 격차: 특수공격"),
            (f"+{gaps['hp']:.1f}", "최소 격차: HP — 화력에 집중"),
        ],
        size=12.5,
        space_after=14,
    )
    add_stat_strip(slide, 8.85, 5.35, 3.38, [("+220.2", "총합 차이"), ("8.1%", "전설 비중"), ("×1.53", "총합 배율")])
    add_notes(
        slide,
        "전설 65마리의 평균 총합은 637.4로 일반 417.2보다 52.8% 높습니다. 여섯 개 능력치 모두에서 앞서며, "
        "특수공격 격차가 53.7로 가장 컸고 HP 격차가 25.6으로 가장 작았습니다. 즉 전설은 내구력보다 화력에서 더 두드러집니다.",
    )


def build_correlation_slide(prs: Presentation, chart: Path) -> None:
    slide = new_content_slide(prs, "07 · CORRELATION", "능력치들은 함께 움직이는가", 9)
    add_image_fit(slide, chart, 0.8, 1.8, 5.7, 4.85)
    add_card(slide, 6.85, 1.85, 5.68, 4.65)
    add_text(slide, 7.15, 2.08, 5.0, 0.35, "상관관계 해석", size=17, bold=True, color=NAVY)
    add_bullets(
        slide,
        7.15,
        2.6,
        5.05,
        3.75,
        [
            ("0.75", "특수공격–총합 상관 (최고)"),
            ("0.74", "공격–총합 상관"),
            ("0.58", "스피드–총합 — 가장 약함"),
            ("0.51", "능력치 간 최대 상관 (방어–특방)"),
            ("특화 경향", "능력치 조합이 특화되는 경향"),
        ],
        size=13,
        space_after=10,
    )
    add_stat_strip(slide, 7.15, 5.35, 5.08, [("0.75", "총합–특공 최대"), ("0.51", "능력치 간 최대"), ("0.02", "최소 상호 상관")])
    add_notes(
        slide,
        "총합과의 상관은 특수공격 0.75, 공격 0.74로 높은 반면 스피드는 0.58로 가장 낮습니다. "
        "능력치끼리의 상관은 최대 0.51에 그쳐, 포켓몬의 능력치가 전반적으로 함께 올라가기보다는 특정 능력치에 특화되는 경향을 보입니다.",
    )


def build_generation_slide(prs: Presentation, chart: Path) -> None:
    slide = new_content_slide(prs, "08 · GENERATION", "세대가 지날수록 강해졌을까", 10)
    add_image_fit(slide, chart, 0.62, 1.8, 7.85, 4.8)
    add_card(slide, 8.7, 1.85, 3.83, 4.65)
    add_text(slide, 9.0, 2.08, 3.3, 0.35, "세대별 특징", size=17, bold=True, color=NAVY)
    add_bullets(
        slide,
        9.0,
        2.58,
        3.25,
        3.7,
        [
            ("459.0", "4세대 평균 — 6개 세대 중 최고"),
            ("418.3", "2세대 평균 — 최저"),
            ("18", "3세대 전설 수 — 최다"),
            ("82–166", "세대별 표본 불균형"),
            ("무추세", "세대가 올라가도 평균은 비슷"),
        ],
        size=12.5,
        space_after=12,
    )
    add_stat_strip(slide, 9.0, 5.35, 3.23, [("82", "최소 표본"), ("166", "최대 표본"), ("459", "최고 평균")])
    add_notes(
        slide,
        "세대별 평균 종족값은 4세대가 459.0으로 가장 높고 2세대가 418.3으로 가장 낮았습니다. "
        "뚜렷한 상승 추세는 없으며, 3세대에 전설이 18마리로 가장 많이 등장했습니다. 세대별 표본 수 차이도 해석에 고려해야 합니다.",
    )


def build_insights_slide(prs: Presentation) -> None:
    slide = new_content_slide(prs, "09 · INSIGHTS", "분석에서 얻은 네 가지 인사이트", 11)
    insights = [
        ("01", "전설은 '다른 리그'", "평균 총합 637.4로 일반보다 52.8% 높고, 6개 능력치 모두 우위를 보인다."),
        ("02", "Dragon이 최강 타입", "평균 541.8로 1위. 최하위 Bug(379.5)와는 162.3 차이가 난다."),
        ("03", "상위권은 전설·메가", "Top10 중 7마리가 전설, 상위 50위의 56%가 전설이다. 일반 3마리는 모두 메가진화다."),
        ("04", "능력치는 특화 설계", "능력치 간 상관은 최대 0.51. 스피드는 총합과 0.58로 약하게 연결된다."),
    ]
    for index, (number, title, body) in enumerate(insights):
        column = index % 2
        row = index // 2
        x = 0.8 + column * 5.98
        y = 1.82 + row * 1.96
        add_card(slide, x, y, 5.75, 1.78)
        add_text(slide, x + 0.3, y + 0.2, 0.6, 0.35, number, size=13, bold=True, color=RED)
        add_text(slide, x + 0.95, y + 0.18, 4.5, 0.38, title, size=16.5, bold=True, color=NAVY)
        add_text(slide, x + 0.95, y + 0.66, 4.45, 0.95, body, size=12.2, color=GRAY, line_spacing=1.25)
    add_rect(slide, 0.8, 5.92, 11.73, 0.78, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    add_text(
        slide,
        1.1,
        6.08,
        11.2,
        0.5,
        "총합은 강함의 대리지표일 뿐 — 실제 전투력은 타입 상성과 기술 구성까지 함께 봐야 한다",
        size=13.5,
        bold=True,
        color=WHITE,
    )
    add_notes(
        slide,
        "네 가지로 정리할 수 있습니다. 전설의 압도적 우위, Dragon의 최강 타입, 상위권의 전설·메가 독식, "
        "그리고 능력치가 서로 독립적으로 특화되는 경향입니다. 다만 총합은 실제 전투력의 대리지표라는 점을 강조하고 싶습니다.",
    )


def build_limitations_slide(prs: Presentation) -> None:
    slide = new_content_slide(prs, "10 · LIMITATIONS", "한계와 다음 단계", 12)
    add_card(slide, 0.8, 1.82, 5.75, 3.55)
    add_text(slide, 1.1, 2.05, 4.8, 0.35, "분석의 한계", size=17, bold=True, color=NAVY)
    add_bullets(
        slide,
        1.1,
        2.55,
        5.15,
        2.65,
        [
            "메가진화·원시회귀가 포함되어 기본형 비교가 왜곡됨",
            "타입 상성·기술·특성 데이터가 없어 실제 전투력은 설명하지 못함",
            "세대별 표본 불균형 (82–166마리)",
            "총합은 6개 능력치의 단순 합",
        ],
        size=12,
        space_after=10,
    )
    add_card(slide, 6.78, 1.82, 5.75, 3.55)
    add_text(slide, 7.08, 2.05, 4.8, 0.35, "다음 단계", size=17, bold=True, color=NAVY)
    add_bullets(
        slide,
        7.08,
        2.55,
        5.18,
        2.65,
        [
            "기본형만 분리해 순수 종족값 비교",
            "타입 상성 기반 배틀 시뮬레이션",
            "능력치 프로필 군집화로 역할 유형 도출",
            "총합 예측 모델 또는 이상치 탐지",
        ],
        size=12,
        space_after=10,
    )
    add_rect(slide, 1.1, 4.2, 5.15, 0.85, PALE_RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    add_text(slide, 1.35, 4.42, 4.7, 0.4, "→ 메가진화 제외 후 재검증", size=12.5, bold=True, color=RED)
    add_rect(slide, 7.08, 4.2, 5.15, 0.85, PALE_YELLOW, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    add_text(slide, 7.33, 4.42, 4.7, 0.4, "→ 타입 상성 결합으로 승률 예측 확장", size=12.5, bold=True, color=NAVY)
    add_rect(slide, 0.8, 5.68, 11.73, 0.92, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    add_text(
        slide,
        1.1,
        5.86,
        11.2,
        0.3,
        "재현 가능성",
        size=12,
        bold=True,
        color=NAVY,
    )
    add_text(
        slide,
        1.1,
        6.17,
        11.2,
        0.3,
        "data/raw/Pokemon.csv · src/analysis.py · src/make_deck.py · src/make_html_deck.py · requirements.txt",
        size=11,
        color=GRAY,
    )
    add_notes(
        slide,
        "한계는 메가진화와 원시회귀가 포함되어 있다는 점, 그리고 타입 상성이나 기술 데이터가 없어 실제 전투력까지는 설명하지 못한다는 점입니다. "
        "다음 단계로 기본형만 분리한 비교와 타입 상성 기반 시뮬레이션을 제안합니다.",
    )


def build_closing_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    add_rect(slide, 10.1, 2.1, 4.0, 4.0, NAVY, line=YELLOW, line_width=2.2, shape=MSO_SHAPE.OVAL)
    add_rect(slide, 10.1, 4.07, 4.0, 0.06, YELLOW)
    add_rect(slide, 11.6, 3.6, 0.95, 0.95, NAVY, line=YELLOW, line_width=2.2, shape=MSO_SHAPE.OVAL)
    add_text(slide, 0.9, 2.3, 7.0, 0.35, "THANK YOU", size=13, bold=True, color=YELLOW)
    add_text(slide, 0.9, 2.85, 8.4, 1.0, "감사합니다", size=48, bold=True, color=WHITE)
    add_text(slide, 0.9, 4.15, 8.2, 0.45, "질문과 토론을 환영합니다", size=19, color=RGBColor(0xC9, 0xD1, 0xDF))
    add_text(
        slide,
        0.9,
        5.55,
        8.0,
        0.7,
        ["데이터: Kaggle — abcsds/pokemon", "분석 및 슬라이드: src/analysis.py · src/make_deck.py · src/make_html_deck.py"],
        size=11,
        color=RGBColor(0x9A, 0xA5, 0xB8),
        line_spacing=1.4,
    )
    add_notes(slide, "이상으로 발표를 마치겠습니다. 감사합니다. 질문 받겠습니다.")


def main() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    setup_matplotlib()
    df = load_data()

    long_types = df.melt(id_vars=["total"], value_vars=["type_1", "type_2"], value_name="type")
    long_types = long_types[long_types["type"] != "None"]
    type_mean = long_types.groupby("type")["total"].mean()
    dragon_mean = type_mean["Dragon"]

    regular = df[~df["legendary"]][STATS].mean()
    legendary = df[df["legendary"]][STATS].mean()
    gaps = (legendary - regular).round(1)
    legendary_total = df[df["legendary"]]["total"].mean()
    regular_total = df[~df["legendary"]]["total"].mean()

    charts = {
        "dist": chart_distribution(df),
        "top10": chart_top10(df),
        "type": chart_type_total(df),
        "legendary": chart_legendary(df),
        "correlation": chart_correlation(df),
        "generation": chart_generation(df),
    }

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    build_title_slide(prs)
    build_agenda_slide(prs)
    build_data_slide(prs, df)
    build_kpi_slide(prs, df, dragon_mean, legendary_total, regular_total)
    build_distribution_slide(prs, charts["dist"])
    build_top10_slide(prs, charts["top10"])
    build_type_slide(prs, charts["type"])
    build_legendary_slide(prs, charts["legendary"], gaps)
    build_correlation_slide(prs, charts["correlation"])
    build_generation_slide(prs, charts["generation"])
    build_insights_slide(prs)
    build_limitations_slide(prs)
    build_closing_slide(prs)

    prs.save(OUTPUT)
    print(f"Saved {OUTPUT} ({len(prs.slides)} slides)")
    print(f"Chart assets: {ASSET_DIR}")


if __name__ == "__main__":
    main()
